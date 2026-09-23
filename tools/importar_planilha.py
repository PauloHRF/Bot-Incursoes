"""Converte a planilha de uma incursão no JSON que o bot lê.

    python tools/importar_planilha.py data/planilhas/incursao_exemplo.xlsx
    python tools/importar_planilha.py minha_incursao.xlsx --saida data/incursoes

Valida tudo antes de gravar: se houver qualquer problema, nada é escrito e os
problemas saem listados. Funciona com planilhas do Excel e do Google Sheets
(baixe como .xlsx em Arquivo > Fazer download > Microsoft Excel).

As salas do meio não vêm daqui: são sorteadas do banco da Organização, que tem
a própria planilha (tools/importar_banco.py).
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.incursoes import (  # noqa: E402
    CAMPOS_DE_CRIATURA,
    ErroDeValidacao,
    criatura_de_colunas,
    de_dict,
)
from src.rules import chave_comparacao  # noqa: E402

CAMPOS_META = (
    "id", "nome", "organizacao", "tamanho", "tier", "lore_inicial", "lore_final",
    "imagem_capa", "recompensa_mes", "pontos_conclusao",
)
CAMPOS_OBJETIVO = (
    "sala_id", "nome", "tipo", "descricao", "imagem", "monstro_nome",
    "monstro_quantidade", "monstro_ca", "monstro_ataque", "monstro_dano", "monstro_hp",
    "recompensa", "pontos_organizacao",
) + tuple(f"monstro_{campo}" for campo in CAMPOS_DE_CRIATURA)
COLUNAS_MONSTROS = ("nome", "quantidade", "ca", "ataque", "dano", "hp")


class ErroDePlanilha(Exception):
    pass


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _aba(wb, nome: str):
    for candidata in wb.sheetnames:
        if chave_comparacao(candidata) == chave_comparacao(nome):
            return wb[candidata]
    raise ErroDePlanilha(
        f"A planilha nao tem a aba '{nome}'. Abas encontradas: {', '.join(wb.sheetnames)}."
    )


def _ler_chave_valor(wb, aba: str, campos: tuple[str, ...]) -> dict[str, Any]:
    ws = _aba(wb, aba)
    lidos: dict[str, Any] = {}
    for campo, valor in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        chave = chave_comparacao(_texto(campo))
        if not chave:
            continue
        for esperado in campos:
            if chave == chave_comparacao(esperado):
                lidos[esperado] = valor
    return lidos


def _escolta(wb) -> list[dict[str, Any]]:
    """A aba 'Monstros': quem luta ao lado do chefe. Opcional."""
    if "Monstros" not in wb.sheetnames:
        return []
    ws = wb["Monstros"]
    cabecalho = [_texto(c.value) for c in ws[1]]
    faltando = [c for c in COLUNAS_MONSTROS if c not in cabecalho]
    if faltando:
        raise ErroDePlanilha(
            f"A aba 'Monstros' esta sem a(s) coluna(s): {', '.join(faltando)}."
        )
    nomes = COLUNAS_MONSTROS + CAMPOS_DE_CRIATURA
    indices = {nome: cabecalho.index(nome) for nome in nomes if nome in cabecalho}
    criaturas = []
    for linha in ws.iter_rows(min_row=2):
        def valor(nome: str):
            indice = indices.get(nome)
            if indice is None or indice >= len(linha):
                return None  # coluna nova que esta planilha ainda nao tem
            return linha[indice].value

        if not _texto(valor("nome")):
            continue
        criaturas.append(criatura_de_colunas(valor))
    return criaturas


def importar(planilha: Path, saida: Path) -> Path:
    if not planilha.exists():
        raise ErroDePlanilha(f"Planilha nao encontrada: {planilha}")

    wb = load_workbook(planilha, data_only=True)
    meta = _ler_chave_valor(wb, "Incursao", CAMPOS_META)
    faltando = [c for c in ("id", "nome", "organizacao") if not _texto(meta.get(c))]
    if faltando:
        raise ErroDePlanilha(
            f"A aba 'Incursao' nao tem: {', '.join(faltando)}. "
            "A coluna A traz o nome do campo e a coluna B o valor."
        )

    bruto = _ler_chave_valor(wb, "Objetivo", CAMPOS_OBJETIVO)
    if not _texto(bruto.get("sala_id")):
        raise ErroDePlanilha("A aba 'Objetivo' precisa de um 'sala_id' na coluna B.")

    objetivo: dict[str, Any] = {
        "id": _texto(bruto.get("sala_id")),
        "nome": _texto(bruto.get("nome")),
        "tipo": _texto(bruto.get("tipo")) or "Combate",
        "descricao": _texto(bruto.get("descricao")),
        "imagem": _texto(bruto.get("imagem")) or None,
        "recompensa": _texto(bruto.get("recompensa")) or None,
        "pontos_organizacao": bruto.get("pontos_organizacao"),
    }
    objetivo["monstros"] = []
    if _texto(bruto.get("monstro_nome")):
        objetivo["monstros"].append(criatura_de_colunas(bruto.get, "monstro_"))
    objetivo["monstros"].extend(_escolta(wb))

    dados = {
        "id": _texto(meta.get("id")),
        "nome": _texto(meta.get("nome")),
        "organizacao": _texto(meta.get("organizacao")),
        "tamanho": _texto(meta.get("tamanho")) or "Média",
        "tier": meta.get("tier"),
        "lore_inicial": _texto(meta.get("lore_inicial")),
        "lore_final": _texto(meta.get("lore_final")) or None,
        "imagem_capa": _texto(meta.get("imagem_capa")) or None,
        "recompensa_mes": meta.get("recompensa_mes"),
        "pontos_conclusao": meta.get("pontos_conclusao"),
        "objetivo": objetivo,
    }

    incursao = de_dict(dados)  # levanta ErroDeValidacao com a lista de problemas

    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / f"{incursao.id}.json"
    destino.write_text(
        json.dumps(incursao.para_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return destino


def main() -> int:
    parser = argparse.ArgumentParser(description="Converte a planilha de incursao em JSON.")
    parser.add_argument("planilha", type=Path, help="Arquivo .xlsx da incursao")
    parser.add_argument(
        "--saida", type=Path, default=Path("data/incursoes"), help="Pasta de destino do JSON"
    )
    args = parser.parse_args()

    try:
        destino = importar(args.planilha, args.saida)
    except ErroDePlanilha as exc:
        print(f"Nao consegui ler a planilha:\n  - {exc}", file=sys.stderr)
        return 1
    except ErroDeValidacao as exc:
        print(f"A planilha tem {len(exc.problemas)} problema(s); nada foi gravado:", file=sys.stderr)
        print(exc, file=sys.stderr)
        return 1

    dados = json.loads(destino.read_text(encoding="utf-8"))
    print(
        f"JSON gravado em: {destino} — {dados['tamanho'].lower()}, "
        f"tier {dados['tier']}, {dados['organizacao']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
