"""Gera a planilha modelo de uma incursão, preenchida com um exemplo.

    python tools/gerar_modelo_planilha.py [destino.xlsx]

A incursão guarda a lore de abertura, a lore de fecho, o tamanho e a sala final.
As salas do meio vêm do banco da Organização (tools/gerar_modelo_banco.py) e são
sorteadas quando a run começa.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.incursoes import ORGANIZACOES, TAMANHOS  # noqa: E402
from src.rules import TIER_MAXIMO  # noqa: E402

FONTE = "Arial"
cab_fonte = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
cab_fill = PatternFill("solid", fgColor="2F3B52")
corpo = Font(name=FONTE, size=10)
corpo_bold = Font(name=FONTE, size=10, bold=True)
calculado = Font(name=FONTE, size=10, color="808080", italic=True)
topo = Alignment(vertical="top", wrap_text=True)

LORE_INICIAL = (
    "Um selo antigo cedeu sob a cidade e algo do outro lado começou a vazar. Três dias depois, "
    "os cães pararam de latir no bairro inteiro e uma rua acordou com as portas trancadas por "
    "dentro e ninguém atrás delas.\n\n"
    "O Vórtice Oculto já sabia do selo — foi um dos seus que o fez, há muito tempo. Agora quer "
    "o Olho de volta antes que a fenda se alargue o bastante para alguém de fora perceber de "
    "quem foi a culpa. Vocês descem hoje."
)

LORE_FINAL = (
    "O Olho para de piscar assim que sai do pedestal, e o silêncio que vem depois é pior que o "
    "zumbido. Lá em cima, os cães voltam a latir.\n\n"
    "A Organização recebe o Olho sem cerimônia, como quem recolhe uma ferramenta emprestada. "
    "Ninguém menciona o que ele guardava, nem por que o selo tinha a marca de um deles. Vocês "
    "recebem o pagamento, e a instrução de não voltar àquela rua."
)

META = [
    ("id", "vortice_cripta", "Identificador curto e único, sem espaço. Vira o nome do arquivo JSON."),
    ("nome", "A Cripta do Vórtice", "Nome exibido aos jogadores."),
    ("organizacao", "Vórtice Oculto", "De qual Organização é a incursão. Define o banco de salas usado no sorteio."),
    ("tamanho", "Média", "Curta (3 salas), Média (5) ou Longa (7), sempre mais o objetivo."),
    ("tier", 4, f"Tier da incursão, de 1 a {TIER_MAXIMO} (um tier a cada 2 níveis). "
     "Só entra quem for deste tier ou de um mais baixo: tier 4 aceita até o nível 8."),
    ("lore_inicial", LORE_INICIAL, "Texto de abertura, postado quando a run começa."),
    ("lore_final", LORE_FINAL, "Epílogo, postado só quando o grupo cumpre o objetivo."),
    ("imagem_capa", "", "URL da imagem de capa. Opcional."),
    ("recompensa_mes", 10, "MEs por participante ao completar o objetivo."),
    ("pontos_conclusao", 10, "Pontos de Organização que a conclusão rende ao servidor."),
]

CABECALHO_MONSTROS = (
    ("nome", 24, "Nome da criatura."),
    ("quantidade", 12, "Quantas iguais. Vazio ou 1 = uma."),
    ("ca", 8, "Classe de Armadura."),
    ("ataque", 10, "Bonus de ataque (ex.: 5)."),
    ("dano", 12, "Dado de dano (ex.: 1d8+2)."),
    ("hp", 8, "Pontos de vida de cada uma."),
)

# A escolta do chefe, na aba 'Monstros'.
ESCOLTA = [
    ("Acólito do Selo", 2, 13, 4, "1d8+2", 18),
]

OBJETIVO = [
    ("sala_id", "OBJ", "Identificador da sala final."),
    ("nome", "O Olho do Vórtice", "Nome da sala final."),
    ("tipo", "Combate", "O objetivo é sempre Combate. Não mude."),
    ("descricao", "A câmara é uma esfera perfeita e o Olho flutua no centro dela, aberto. O que o "
     "guarda não tem nome porque nada que o viu voltou para dar um.", "Texto do embed da sala final."),
    ("imagem", "", "Caminho em assets/ ou URL. Opcional."),
    ("monstro_nome", "Guardião do Selo", "Nome do chefe."),
    ("monstro_quantidade", 1, "Quantos chefes iguais. Vazio ou 1 = um. O total da sala, "
     "contando a aba 'Monstros', nao passa de 6."),
    ("monstro_ca", 16, "Classe de Armadura do chefe."),
    ("monstro_ataque", 7, "Bônus de ataque do chefe."),
    ("monstro_dano", "2d8+4", "Dado de dano do chefe (ex.: 2d8+4)."),
    ("monstro_hp", 58, "HP do chefe. Calibre com tools/simular_combate.py."),
    ("recompensa", "10 MEs por participante + pontos com o Vórtice Oculto", "Texto da recompensa."),
    ("pontos_organizacao", 0, "Pontos extras da sala final, além de pontos_conclusao."),
]

LEIA_ME = [
    ("Planilha de incursão — Incursões 2.0", True),
    ("", False),
    ("O que fica aqui", True),
    ("A incursao guarda a lore, o tamanho e a sala final. As salas do meio NAO ficam aqui:", False),
    ("elas sao sorteadas do banco da Organizacao a cada run, entao duas runs da mesma", False),
    ("incursao percorrem caminhos diferentes.", False),
    ("", False),
    ("O banco de salas fica em outra planilha:", False),
    ("   python tools/gerar_modelo_banco.py \"Vortice Oculto\"", False),
    ("   python tools/importar_banco.py data/planilhas/banco_vortice_oculto.xlsx", False),
    ("", False),
    ("Tamanhos", True),
    ("Curta = 3 salas, Media = 5, Longa = 7 — sempre mais a sala de Objetivo no fim.", False),
    ("", False),
    ("Como usar", True),
    ("1. Preencha a aba 'Incursao' e a aba 'Objetivo'.", False),
    ("2. Converta:  python tools/importar_planilha.py <esta planilha>", False),
    ("3. No Discord, /incursao recarregar para o bot reler sem reiniciar.", False),
    ("", False),
    ("Lore", True),
    ("lore_inicial abre a run, assim que o grupo fecha. lore_final so aparece se o grupo", False),
    ("cumprir o objetivo — e o epilogo da historia, nao um resumo do que aconteceu.", False),
]


def _aba_chave_valor(wb, nome: str, linhas: list[tuple], titulo_valor: str):
    ws = wb.create_sheet(nome)
    for i, (titulo, largura) in enumerate((("campo", 20), (titulo_valor, 78), ("o que é", 58)), start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = cab_fonte
        c.fill = cab_fill
        c.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    for i, (campo, valor, ajuda) in enumerate(linhas, start=2):
        ws.cell(row=i, column=1, value=campo).font = corpo_bold
        v = ws.cell(row=i, column=2, value=valor)
        v.font = corpo
        v.alignment = topo
        a = ws.cell(row=i, column=3, value=ajuda)
        a.font = calculado
        a.alignment = topo
        ws.row_dimensions[i].height = 70 if len(str(valor)) > 120 else 32
    return ws


def gerar(destino: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Leia-me")
    ws.column_dimensions["A"].width = 100
    for i, (texto, destaque) in enumerate(LEIA_ME, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.font = corpo_bold if destaque else corpo
    ws.sheet_view.showGridLines = False

    ws = _aba_chave_valor(wb, "Incursao", META, "valor")
    dv_org = DataValidation(type="list", formula1=f'"{",".join(ORGANIZACOES)}"', allow_blank=False)
    ws.add_data_validation(dv_org)
    dv_org.add("B4")
    dv_tam = DataValidation(type="list", formula1=f'"{",".join(TAMANHOS)}"', allow_blank=False)
    ws.add_data_validation(dv_tam)
    dv_tam.add("B5")
    dv_tier = DataValidation(
        type="whole", operator="between", formula1=1, formula2=TIER_MAXIMO, allow_blank=True
    )
    ws.add_data_validation(dv_tier)
    dv_tier.add("B6")
    ws.cell(row=1, column=2).comment = Comment(
        "Uma linha por campo. Nao apague nem renomeie os campos da coluna A.",
        "Incursoes 2.0",
        height=80,
        width=260,
    )

    _aba_chave_valor(wb, "Objetivo", OBJETIVO, "valor")

    ws = wb.create_sheet("Monstros")
    for i, (titulo, largura, ajuda) in enumerate(CABECALHO_MONSTROS, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = cab_fonte
        c.fill = cab_fill
        c.comment = Comment(ajuda, "Incursoes 2.0", height=80, width=260)
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.cell(row=1, column=1).comment = Comment(
        "Criaturas que lutam AO LADO do chefe. O que estiver aqui se soma ao monstro_*"
        " da aba Objetivo, ate 6 no total. Deixe vazia para um chefe sozinho.",
        "Incursoes 2.0",
        height=110,
        width=300,
    )
    for i, linha in enumerate(ESCOLTA, start=2):
        for col, valor in enumerate(linha, start=1):
            ws.cell(row=i, column=col, value=valor).font = corpo
    ws.freeze_panes = "A2"

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


if __name__ == "__main__":
    alvo = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data/planilhas/incursao_exemplo.xlsx")
    print("planilha gerada:", gerar(alvo))
