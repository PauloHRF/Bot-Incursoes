"""Converte a planilha do banco de salas de uma Organização no JSON que o bot lê.

    python tools/importar_banco.py data/planilhas/banco_vortice_oculto.xlsx

Valida tudo antes de gravar: se houver qualquer problema, nada é escrito.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.incursoes import (  # noqa: E402
    ErroDeValidacao,
    arquivo_da_organizacao,
    banco_de_dict,
)
from src.rules import chave_comparacao  # noqa: E402

COLUNAS = (
    "sala_id", "nome", "tipo", "dificuldade", "cd", "alvo_progresso", "pericias",
    "descricao", "imagem", "monstro_nome", "monstro_quantidade", "monstro_ca",
    "monstro_ataque", "monstro_dano", "monstro_hp", "recompensa", "pontos_organizacao",
)
COLUNAS_MONSTROS = ("sala_id", "nome", "quantidade", "ca", "ataque", "dano", "hp")


class ErroDePlanilha(Exception):
    pass


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _aba(wb, nome: str):
    for candidata in wb.sheetnames:
        if chave_comparacao(candidata) == chave_comparacao(nome):
            return wb[candidata]
    raise ErroDePlanilha(
        f"A planilha nao tem a aba '{nome}'. Abas encontradas: {', '.join(wb.sheetnames)}."
    )


def ler_dificuldades(wb) -> dict[str, tuple[int | None, int | None]]:
    ws = _aba(wb, "Dificuldades")
    tabela = {}
    for nome, cd, alvo in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if not _texto(nome):
            continue
        tabela[chave_comparacao(_texto(nome))] = (
            int(cd) if cd is not None else None,
            int(alvo) if alvo is not None else None,
        )
    if not tabela:
        raise ErroDePlanilha("A aba 'Dificuldades' esta vazia.")
    return tabela


def ler_organizacao(wb) -> str:
    ws = _aba(wb, "Banco")
    for campo, valor in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if chave_comparacao(_texto(campo)) == chave_comparacao("organizacao"):
            if _texto(valor):
                return _texto(valor)
    raise ErroDePlanilha(
        "A aba 'Banco' precisa de uma linha com 'organizacao' na coluna A e o nome na coluna B."
    )


def ler_salas(wb, dificuldades) -> list[dict[str, Any]]:
    ws = _aba(wb, "Salas")
    cabecalho = [chave_comparacao(_texto(c.value)) for c in ws[1]]
    try:
        col = {nome: cabecalho.index(chave_comparacao(nome)) for nome in COLUNAS}
    except ValueError as exc:
        raise ErroDePlanilha(
            f"A aba 'Salas' esta sem uma coluna obrigatoria ({exc}). "
            "Regenere o modelo com tools/gerar_modelo_banco.py e compare os cabecalhos."
        ) from exc

    salas = []
    for numero, celulas in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def valor(nome):
            indice = col[nome]
            return celulas[indice] if indice < len(celulas) else None

        sala_id = _texto(valor("sala_id"))
        if not sala_id:
            continue

        dif = _texto(valor("dificuldade"))
        cd, alvo = valor("cd"), valor("alvo_progresso")
        # A planilha calcula cd/alvo por formula; sem cache, resolvemos aqui.
        if dif and chave_comparacao(dif) not in dificuldades:
            raise ErroDePlanilha(
                f"Linha {numero}: dificuldade '{dif}' nao esta na aba 'Dificuldades'."
            )
        if dif and (cd is None or isinstance(cd, str)):
            cd = dificuldades[chave_comparacao(dif)][0]
        if dif and (alvo is None or isinstance(alvo, str)):
            alvo = dificuldades[chave_comparacao(dif)][1]

        sala: dict[str, Any] = {
            "id": sala_id,
            "nome": _texto(valor("nome")),
            "tipo": _texto(valor("tipo")),
            "descricao": _texto(valor("descricao")),
            "dificuldade": dif or None,
            "cd": cd,
            "alvo_progresso": alvo,
            "pericias": [
                p.strip()
                for p in _texto(valor("pericias")).replace(",", ";").split(";")
                if p.strip()
            ],
            "imagem": _texto(valor("imagem")) or None,
            "recompensa": _texto(valor("recompensa")) or None,
            "pontos_organizacao": valor("pontos_organizacao"),
        }
        sala["monstros"] = []
        if _texto(valor("monstro_nome")):
            sala["monstros"].append(
                {
                    "nome": _texto(valor("monstro_nome")),
                    "quantidade": valor("monstro_quantidade"),
                    "ca": valor("monstro_ca"),
                    "ataque": valor("monstro_ataque"),
                    "dano": _texto(valor("monstro_dano")),
                    "hp": valor("monstro_hp"),
                }
            )
        salas.append(sala)
    return salas


def _monstros_extras(wb) -> dict[str, list[dict[str, Any]]]:
    """A aba 'Monstros': criaturas a mais de uma sala, para bandos mistos.

    A aba e opcional; sem ela, cada sala tem so a criatura da propria linha.
    """
    try:
        ws = _aba(wb, "Monstros")
    except ErroDePlanilha:
        return {}
    cabecalho = [_texto(c.value) for c in ws[1]]
    indices = {nome: cabecalho.index(nome) for nome in COLUNAS_MONSTROS if nome in cabecalho}
    faltando = [c for c in COLUNAS_MONSTROS if c not in indices]
    if faltando:
        raise ErroDePlanilha(
            f"A aba 'Monstros' esta sem a(s) coluna(s): {', '.join(faltando)}."
        )
    por_sala: dict[str, list[dict[str, Any]]] = {}
    for linha in ws.iter_rows(min_row=2):
        def valor(nome: str):
            celula = linha[indices[nome]]
            return celula.value

        sala_id = _texto(valor("sala_id"))
        nome = _texto(valor("nome"))
        if not sala_id and not nome:
            continue
        if not sala_id:
            raise ErroDePlanilha(f"A aba 'Monstros' tem a criatura '{nome}' sem sala_id.")
        por_sala.setdefault(sala_id, []).append(
            {
                "nome": nome,
                "quantidade": valor("quantidade"),
                "ca": valor("ca"),
                "ataque": valor("ataque"),
                "dano": _texto(valor("dano")),
                "hp": valor("hp"),
            }
        )
    return por_sala


def importar(planilha: Path, saida: Path) -> Path:
    if not planilha.exists():
        raise ErroDePlanilha(f"Planilha nao encontrada: {planilha}")

    wb = load_workbook(planilha, data_only=True)
    dificuldades = ler_dificuldades(wb)
    salas = ler_salas(wb, dificuldades)
    extras = _monstros_extras(wb)
    conhecidas = {s["id"] for s in salas}
    for sala_id, criaturas in extras.items():
        if sala_id not in conhecidas:
            raise ErroDePlanilha(
                f"A aba 'Monstros' cita a sala '{sala_id}', que nao existe na aba 'Salas'."
            )
        for sala in salas:
            if sala["id"] == sala_id:
                sala["monstros"].extend(criaturas)
    dados = {"organizacao": ler_organizacao(wb), "salas": salas}
    banco = banco_de_dict(dados)

    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / f"{arquivo_da_organizacao(banco.organizacao)}.json"
    destino.write_text(
        json.dumps(banco.para_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return destino


def main() -> int:
    parser = argparse.ArgumentParser(description="Converte a planilha de banco de salas em JSON.")
    parser.add_argument("planilha", type=Path, help="Arquivo .xlsx do banco")
    parser.add_argument("--saida", type=Path, default=Path("data/bancos"))
    args = parser.parse_args()

    try:
        destino = importar(args.planilha, args.saida)
    except ErroDePlanilha as exc:
        print(f"Nao consegui ler a planilha:\n  - {exc}", file=sys.stderr)
        return 1
    except ErroDeValidacao as exc:
        print(f"O banco tem {len(exc.problemas)} problema(s); nada foi gravado:", file=sys.stderr)
        print(exc, file=sys.stderr)
        return 1

    dados = json.loads(destino.read_text(encoding="utf-8"))
    print(f"JSON gravado em: {destino} ({len(dados['salas'])} salas)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
